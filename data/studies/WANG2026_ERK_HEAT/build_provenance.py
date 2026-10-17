"""Record what PXD065668 / PXD065706 hold, and why they are registered but not reanalyzed.

Wang et al. 2026 (Commun Biol 9:844, CC BY-NC-ND 4.0) injected C. gigas with
the ERK inhibitor SCH772984 or sterile seawater, then sampled gill before (0 h)
and after 12 h at 37 C: 4 groups x 3 pools of 3 oysters, DIA on an Orbitrap
Exploris 480, quantified in Spectronaut 18. PRIDE PXD065668 (proteome) and
PXD065706 (phosphoproteome) each hold 12 .raw files and one Spectronaut .sne
archive (proprietary; no open export). The paper's Supplementary Data 1 and 2
give per-pool intensities only for proteins/sites already called
differentially abundant.

    python -I data/studies/WANG2026_ERK_HEAT/build_provenance.py --download STAGING_DIR

fetches both PRIDE projects' records, the first KB of each proteome .raw file
(acquisition start) and the paper's supplementary files (Europe PMC), and
writes run_order.tsv and pride_provenance.json. The supplement is
CC BY-NC-ND, so it is checksummed and summarized, not committed.
"""
from __future__ import annotations

import argparse
import hashlib
import io
import json
import struct
import sys
import urllib.request
import zipfile
from collections import Counter
from datetime import date, datetime, timedelta
from pathlib import Path

import openpyxl
import pandas as pd

STUDY_ID = "WANG2026_ERK_HEAT"
PROJECTS = {"proteome": "PXD065668", "phosphoproteome": "PXD065706"}
PRIDE_API = "https://www.ebi.ac.uk/pride/ws/archive/v3/projects/{}"
PMCID = "PMC13280272"
SUPP_URL = f"https://www.ebi.ac.uk/europepmc/webservices/rest/{PMCID}/supplementaryFiles"
SUPP_TABLES = {"proteome": "42003_2026_10054_MOESM2_ESM.xlsx", "phosphoproteome": "42003_2026_10054_MOESM3_ESM.xlsx"}

REPO_ROOT = Path(__file__).resolve().parents[3]
OUT = REPO_ROOT / "data" / "studies" / STUDY_ID


def fetch(url: str, byte_range: str | None = None) -> bytes:
    headers = {"User-Agent": "AREE-curation"}
    if byte_range:
        headers["Range"] = f"bytes={byte_range}"
    with urllib.request.urlopen(urllib.request.Request(url, headers=headers), timeout=600) as resp:
        return resp.read()


def raw_start_time(header: bytes) -> str:
    ticks = struct.unpack_from("<Q", header, 40)[0]
    return (datetime(1601, 1, 1) + timedelta(microseconds=ticks / 10)).isoformat(timespec="seconds")


def ftp_url(file_record: dict) -> str:
    loc = next(x["value"] for x in file_record["publicFileLocations"] if x["value"].startswith("ftp://"))
    return loc.replace("ftp://", "https://", 1)


def summarize_table(xlsx: bytes) -> dict:
    """Rows per sheet and the regulation calls they hold (all significant, none 'not changed')."""
    wb = openpyxl.load_workbook(io.BytesIO(xlsx), read_only=True)
    out = {"caption": None, "sheets": {}}
    for name in wb.sheetnames:
        rows = list(wb[name].iter_rows(values_only=True))
        start = 1 if rows and rows[0] and str(rows[0][0]).startswith("Supplementary Data") else 0
        if start:
            out["caption"] = str(rows[0][0])
        header = [str(h) for h in rows[start]]
        body = [r for r in rows[start + 1:] if r and r[0] is not None]
        reg = next(i for i, h in enumerate(header) if "Regulation" in h)
        intensity_cols = [h for h in header if h.startswith(("SW_", "ERK_", "Intensity SW_", "Intensity ERK_"))
                          and "Ratio" not in h and "P value" not in h and "/" not in h]
        out["sheets"][name] = {
            "n_rows": len(body),
            "regulation": dict(Counter(str(r[reg]) for r in body)),
            "per_pool_value_columns": intensity_cols,
        }
    return out


def main() -> int:
    ap = argparse.ArgumentParser(description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
    ap.add_argument("--download", type=Path, required=True, help="staging directory for fetched files")
    args = ap.parse_args()
    stage = args.download
    stage.mkdir(parents=True, exist_ok=True)

    projects, runs = {}, []
    for dataset, acc in PROJECTS.items():
        project = json.loads(fetch(PRIDE_API.format(acc)))
        files = json.loads(fetch(PRIDE_API.format(acc) + "/files/all"))
        (stage / f"{acc}_files.json").write_text(json.dumps(files))
        raws = sorted((f for f in files if f["fileCategory"]["value"] == "RAW"), key=lambda f: f["fileName"])
        groups = Counter("_".join(f["fileName"].removesuffix(".raw").split("_")[1:3]) for f in raws)
        if set(groups.values()) != {3} or len(groups) != 4:
            sys.exit(f"{acc}: expected 4 groups x 3 pools, got {dict(groups)}")
        projects[dataset] = {
            "accession": acc,
            "title": project["title"],
            "license": project["license"],
            "submission_type": project["submissionType"],
            "publication_date": project["publicationDate"],
            "instrument": [i["name"] for i in project["instruments"]],
            "files_by_category": dict(sorted(Counter(f["fileCategory"]["value"] for f in files).items())),
            "search_files": [f["fileName"] for f in files if f["fileCategory"]["value"] == "SEARCH"],
            "pools_per_group": dict(sorted(groups.items())),
        }
        if dataset == "proteome":
            for f in raws:
                _, inhibitor, hours, pool = f["fileName"].removesuffix(".raw").split("_")
                runs.append({"raw_file": f["fileName"], "injection": {"SW": "seawater", "ERK": "SCH772984"}[inhibitor],
                             "heat_h": int(hours), "pool": int(pool),
                             "acquired": raw_start_time(fetch(ftp_url(f), "0-1023"))})
    runs.sort(key=lambda r: r["acquired"])
    for i, r in enumerate(runs, start=1):
        r["run_order"] = i
    pd.DataFrame(runs).to_csv(OUT / "run_order.tsv", sep="\t", index=False)

    supp_zip = stage / "supplementaryFiles.zip"
    if not supp_zip.exists():
        supp_zip.write_bytes(fetch(SUPP_URL))
    tables = {}
    with zipfile.ZipFile(supp_zip) as zf:
        for dataset, name in SUPP_TABLES.items():
            raw = zf.read(name)
            tables[dataset] = {"file": name, "sha256": hashlib.sha256(raw).hexdigest(), **summarize_table(raw)}

    provenance = {
        "study_id": STUDY_ID,
        "date_retrieved": date.today().isoformat(),
        "generated_by": f"data/studies/{STUDY_ID}/build_provenance.py",
        "source_urls": [PRIDE_API.format(a) for a in PROJECTS.values()] + [SUPP_URL],
        "paper_license": "CC BY-NC-ND 4.0 (supplement fetched and checksummed, not committed)",
        "pride": projects,
        "proteome_acquisition_order": [f"{r['injection']}_{r['heat_h']}h_{r['pool']}" for r in runs],
        "supplementary_tables": tables,
        "outputs_sha256": {f"data/studies/{STUDY_ID}/run_order.tsv":
                           hashlib.sha256((OUT / "run_order.tsv").read_bytes()).hexdigest()},
    }
    (OUT / "pride_provenance.json").write_text(json.dumps(provenance, indent=2, ensure_ascii=False) + "\n")
    print(json.dumps({"acquisition": provenance["proteome_acquisition_order"],
                      "proteome_sheets": {k: v["n_rows"] for k, v in tables["proteome"]["sheets"].items()}}, indent=1))
    return 0


if __name__ == "__main__":
    sys.exit(main())
