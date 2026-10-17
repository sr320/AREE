"""Build the proteomics workflow inputs for TIMMINS2014_OA_GILL (PXD000835).

The paper (Timmins-Schiffman et al. 2014, BMC Genomics 15:951, CC BY 4.0)
publishes the per-oyster NSAF table the authors analysed (Additional file 4,
Table S3: 1,616 proteins x 16 oysters) and the ProteinProphet output of every
technical replicate (Additional file 3, Table S2, one sheet per MS run named
<oyster>_<rep>). Table S3 labels oysters by treatment and replicate ("2800 1");
Table S2 and PRIDE label them by oyster number (2, 5, ..., 251). The authors'
SQLShare workflow on figshare (10.6084/m9.figshare.894323) states which oyster
numbers belong to which treatment. This script ties the two labellings
together by data, not by column order: each Table S3 column must correlate
best with the mean spectral counts of exactly one oyster's three Table S2
sheets, and that oyster must be in the treatment the column names.

    python -I data/studies/TIMMINS2014_OA_GILL/build_inputs.py --download STAGING_DIR

fetches the paper's supplementary files (Europe PMC), the figshare workflow and
the first KB of each PRIDE .raw file (for acquisition times), and writes

  nsaf_matrix.tsv                  protein x oyster NSAF (Table S3, columns renamed to oyster IDs)
  samplesheet.tsv                  one row per oyster
  samplesheet_<comparison>.tsv     sample_id/group sheets for the workflow
  run_order.tsv                    the 48 MS runs in acquisition order
  inputs_provenance.json

The NSAF values are copied unchanged. The supplement is CC BY 4.0, so the
derived matrix is committed with attribution in the study YAML.
"""
from __future__ import annotations

import argparse
import hashlib
import io
import json
import re
import struct
import sys
import urllib.request
import zipfile
from datetime import date, datetime, timedelta
from pathlib import Path

import numpy as np
import openpyxl
import pandas as pd

STUDY_ID = "TIMMINS2014_OA_GILL"
ACCESSION = "PXD000835"
PMCID = "PMC4531390"
SUPP_URL = f"https://www.ebi.ac.uk/europepmc/webservices/rest/{PMCID}/supplementaryFiles"
TABLE_S2 = "12864_2014_7071_MOESM3_ESM.xlsx"
TABLE_S3 = "12864_2014_7071_MOESM4_ESM.txt"
FIGSHARE_API = "https://api.figshare.com/v2/articles/894323"
PRIDE_API = f"https://www.ebi.ac.uk/pride/ws/archive/v3/projects/{ACCESSION}/files/all"
PRIDE_FTP = "https://ftp.pride.ebi.ac.uk/pride/data/archive/2014/05/PXD000835"

REPO_ROOT = Path(__file__).resolve().parents[3]
OUT = REPO_ROOT / "data" / "studies" / STUDY_ID

# figshare NSAFworkflow.rtf: "Oysters 2-11 correspond to the 2800 uatm pCO2
# treatment, 26-35 are 2800 uatm + mechanical stress, 221-230 are 400 uatm
# pCO2, and 242-251 are 400 uatm + mechanical stress."
GROUPS = {
    ("2800", False): [2, 5, 8, 11],
    ("2800", True): [26, 29, 32, 35],
    ("400", False): [221, 224, 227, 230],
    ("400", True): [242, 245, 248, 251],
}
WORKFLOW_SENTENCE = "Oysters 2-11 correspond to the 2800 atm pCO2 treatment, 26-35 are 2800 atm + mechanical stress , 221-230 are 400 atm pCO2, and 242-251 are 400 atm + mechanical stress"

COMPARISONS = {
    # comparison_id: (control (pco2, mech), treatment (pco2, mech))
    "pco2_2800_vs_400": (("400", False), ("2800", False)),
    "mech_stress_vs_none_at_400": (("400", False), ("400", True)),
    "mech_stress_vs_none_at_2800": (("2800", False), ("2800", True)),
}


def fetch(url: str, byte_range: str | None = None) -> bytes:
    headers = {"User-Agent": "AREE-curation"}
    if byte_range:
        headers["Range"] = f"bytes={byte_range}"
    with urllib.request.urlopen(urllib.request.Request(url, headers=headers), timeout=600) as resp:
        return resp.read()


def sha256(data: bytes) -> str:
    return hashlib.sha256(data).hexdigest()


def rtf_text(raw: bytes) -> str:
    """Plain text of an RTF file: drop hex escapes, control words and braces."""
    text = raw.decode("latin-1")
    text = re.sub(r"\\'[0-9a-f]{2}", " ", text)
    text = re.sub(r"\\[a-zA-Z]+-?\d* ?", "", text)
    return re.sub(r"\s+", " ", re.sub(r"[{}]", "", text))


def table_s3_column(pco2: str, mech: bool, k: int) -> str:
    return f"{pco2} {'MechS' if mech else ''}{k}"


def oyster_spectral_counts(xlsx: bytes) -> dict[int, pd.Series]:
    """Mean 'tot indep spectra' per protein over each oyster's three runs (Table S2)."""
    wb = openpyxl.load_workbook(io.BytesIO(xlsx), read_only=True)
    per_run = {}
    for name in wb.sheetnames:
        rows = wb[name].iter_rows(values_only=True)
        header = [str(h) for h in next(rows)]
        i_prot, i_spc = header.index("protein"), header.index("tot indep spectra")
        counts = {}
        for r in rows:
            if r[i_prot] is None:
                continue
            for prot in str(r[i_prot]).split(","):
                counts[prot] = max(counts.get(prot, 0), float(r[i_spc] or 0))
        per_run[name] = pd.Series(counts)
    oysters = sorted({int(n.split("_")[0]) for n in per_run})
    return {o: pd.concat([per_run[n] for n in per_run if int(n.split("_")[0]) == o], axis=1).fillna(0).mean(axis=1)
            for o in oysters}


def raw_start_time(header: bytes) -> str:
    # Thermo .raw file header: FILETIME of acquisition start at byte offset 40.
    ticks = struct.unpack_from("<Q", header, 40)[0]
    return (datetime(1601, 1, 1) + timedelta(microseconds=ticks / 10)).isoformat(timespec="seconds")


def main() -> int:
    ap = argparse.ArgumentParser(description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
    ap.add_argument("--download", type=Path, required=True, help="staging directory for fetched files")
    args = ap.parse_args()
    stage = args.download
    stage.mkdir(parents=True, exist_ok=True)

    supp_zip = stage / "supplementaryFiles.zip"
    if not supp_zip.exists():
        supp_zip.write_bytes(fetch(SUPP_URL))
    with zipfile.ZipFile(supp_zip) as zf:
        s2_raw, s3_raw = zf.read(TABLE_S2), zf.read(TABLE_S3)

    article = json.loads(fetch(FIGSHARE_API))
    (wf_file,) = article["files"]
    wf_raw = fetch(wf_file["download_url"])
    if hashlib.md5(wf_raw).hexdigest() != wf_file["computed_md5"]:
        sys.exit("figshare NSAFworkflow.rtf md5 does not match the figshare API")
    if WORKFLOW_SENTENCE not in rtf_text(wf_raw):
        sys.exit("figshare workflow no longer states the oyster-number -> treatment mapping this script encodes")

    # Table S3: classic-Mac line endings, per-oyster NSAF in columns 2-17.
    s3 = pd.read_csv(io.StringIO(s3_raw.decode("latin-1").replace("\r", "\n")), sep="\t", dtype=str)
    s3 = s3.rename(columns={s3.columns[0]: "protein_accession"})
    if len(s3) != 1616 or s3["protein_accession"].duplicated().any():
        sys.exit(f"Table S3: expected 1,616 unique proteins, got {len(s3)}")

    # Tie each Table S3 column to an oyster by correlation with Table S2 spectral
    # counts. Raw abundances correlate ~0.95 between any two oysters (the same
    # abundant proteins dominate every gill), which cannot tell oysters apart; the
    # per-protein deviation from the 16-oyster mean is what is oyster-specific.
    spc = oyster_spectral_counts(s2_raw)
    if sorted(spc) != sorted(o for g in GROUPS.values() for o in g):
        sys.exit(f"Table S2 oysters {sorted(spc)} differ from the figshare grouping")
    s3_cols = [table_s3_column(p, m, k) for (p, m), nums in GROUPS.items() for k in range(1, len(nums) + 1)]
    nsaf_all = s3.set_index("protein_accession")[s3_cols].apply(pd.to_numeric)
    spc_all = pd.DataFrame({o: spc[o].reindex(nsaf_all.index).fillna(0) for o in spc})
    floor = nsaf_all[nsaf_all > 0].min().min() / 2

    def deviations(m: pd.DataFrame) -> pd.DataFrame:
        logged = np.log2(m + floor) if m is nsaf_all else np.log2(m + 1)
        return logged.sub(logged.mean(axis=1), axis=0)

    dev_nsaf, dev_spc = deviations(nsaf_all), deviations(spc_all)
    oysters, columns = [], {}
    for (pco2, mech), numbers in GROUPS.items():
        for k, expected in enumerate(numbers, start=1):
            col = table_s3_column(pco2, mech, k)
            cors = {o: np.corrcoef(dev_nsaf[col], dev_spc[o])[0, 1] for o in spc}
            best = max(cors, key=cors.get)
            runner_up = sorted(cors.values())[-2]
            if best != expected:
                sys.exit(f"Table S3 column {col!r} matches oyster {best}, not {expected}")
            columns[col] = f"oy{expected:03d}"
            oysters.append({
                "sample_id": f"oy{expected:03d}", "oyster_number": expected, "table_s3_column": col,
                "pco2_uatm": pco2, "mechanical_stress": "yes" if mech else "no",
                "pearson_check_r": round(cors[best], 3), "next_best_r": round(runner_up, 3),
            })

    nsaf = s3[["protein_accession", *columns]].rename(columns=columns)
    nsaf.to_csv(OUT / "nsaf_matrix.tsv", sep="\t", index=False)
    pd.DataFrame(oysters).to_csv(OUT / "samplesheet.tsv", sep="\t", index=False)
    by_group = {(o["pco2_uatm"], o["mechanical_stress"] == "yes"): [] for o in oysters}
    for o in oysters:
        by_group[(o["pco2_uatm"], o["mechanical_stress"] == "yes")].append(o["sample_id"])
    for cid, (ctl, trt) in COMPARISONS.items():
        rows = [(s, "control") for s in by_group[ctl]] + [(s, "treatment") for s in by_group[trt]]
        pd.DataFrame(rows, columns=["sample_id", "group"]).to_csv(OUT / f"samplesheet_{cid}.tsv", sep="\t", index=False)

    # Acquisition order of the 48 runs, from each .raw header.
    files = json.loads(fetch(PRIDE_API))
    runs = []
    for f in sorted(f["fileName"] for f in files if f["fileName"].endswith(".raw")):
        _, batch, oyster, rep = f.removesuffix(".raw").split("_")
        group = next(g for g, nums in GROUPS.items() if int(oyster) in nums)
        runs.append({"raw_file": f, "sample_id": f"oy{int(oyster):03d}", "sample_label": batch,
                     "technical_replicate": int(rep), "pco2_uatm": group[0],
                     "mechanical_stress": "yes" if group[1] else "no",
                     "acquired": raw_start_time(fetch(f"{PRIDE_FTP}/{f}", "0-1023"))})
    runs.sort(key=lambda r: r["acquired"])
    for i, r in enumerate(runs, start=1):
        r["run_order"] = i
    pd.DataFrame(runs).to_csv(OUT / "run_order.tsv", sep="\t", index=False)
    rep1 = [r["pco2_uatm"] for r in runs if r["technical_replicate"] == 1]

    outputs = sorted(p.name for p in OUT.glob("*.tsv"))
    provenance = {
        "study_id": STUDY_ID,
        "accession": ACCESSION,
        "date_retrieved": date.today().isoformat(),
        "generated_by": f"data/studies/{STUDY_ID}/build_inputs.py",
        "source_urls": [SUPP_URL, FIGSHARE_API, PRIDE_API],
        "license": "CC BY 4.0 (paper and supplementary files; figshare workflow); PRIDE: EBI terms of use",
        "input_sha256": {
            TABLE_S2: sha256(s2_raw),
            TABLE_S3: sha256(s3_raw),
            wf_file["name"]: sha256(wf_raw),
        },
        "n_proteins": int(len(nsaf)),
        "n_oysters": len(oysters),
        "oysters_by_group": {f"{p} uatm{' + mech' if m else ''}": n for (p, m), n in GROUPS.items()},
        "column_mapping_check": ("each Table S3 column's per-protein log deviations from the 16-oyster mean "
                                 "correlate best with exactly the expected oyster's Table S2 spectral counts"),
        "min_mapping_margin_r": round(min(o["pearson_check_r"] - o["next_best_r"] for o in oysters), 3),
        "n_runs": len(runs),
        "acquisition_window": [runs[0]["acquired"], runs[-1]["acquired"]],
        "technical_replicate_1_order": rep1,
        "comparisons": {cid: {"control": f"{c[0]} uatm{' + mech' if c[1] else ''}",
                              "treatment": f"{t[0]} uatm{' + mech' if t[1] else ''}"}
                        for cid, (c, t) in COMPARISONS.items()},
        "outputs_sha256": {f"data/studies/{STUDY_ID}/{n}": sha256((OUT / n).read_bytes()) for n in outputs},
    }
    (OUT / "inputs_provenance.json").write_text(json.dumps(provenance, indent=2) + "\n")
    print(f"wrote {len(outputs)} tables and inputs_provenance.json; mapping margin r >= {provenance['min_mapping_margin_r']}")
    return 0


if __name__ == "__main__":
    sys.exit(main())
