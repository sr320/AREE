"""Build the metabolomics workflow inputs for DUPERRET2025_POMS_STARVE from Table S12.

MetaboLights MTBLS13313 has no abundance values (its MAFs are empty), but the
paper's Table S12 (mmc13.xlsx, "Raw metabolomic data") holds per-sample peak
areas for all four targeted platforms. The paper is CC BY-NC-ND 4.0, so the
table is NOT committed: this script fetches the supplement from Europe PMC,
checks its SHA-256, and writes the workflow inputs to a staging directory.
What is committed is this script and the provenance JSON it writes.

    python -I data/studies/DUPERRET2025_POMS_STARVE/build_feature_table.py --download STAGING_DIR

writes, in STAGING_DIR:

  mmc13.xlsx                                the supplement, checksum-verified
  DUPERRET2025_POMS_STARVE_features.tsv     feature x sample peak areas, 48 samples
  DUPERRET2025_POMS_STARVE_starved_vs_fed_0hpc_samplesheet.tsv
                                            the 12 samples of the 0 hpc contrast

and data/studies/DUPERRET2025_POMS_STARVE/table_s12_provenance.json.

The reshaping is mechanical. Values are copied unchanged, including the 0.0001
and 0 below-detection placeholders (blank cells stay missing): the workflow's --metabolomics.detection_limit
turns them into missing values. Internal standards are the rows the authors
filled green ("In green are the internal standard"); the fill and a list of
the expected names must agree, or the script stops.
"""
from __future__ import annotations

import argparse
import hashlib
import io
import json
import sys
import urllib.request
import zipfile
from datetime import date
from pathlib import Path

import openpyxl
import pandas as pd

STUDY_ID = "DUPERRET2025_POMS_STARVE"
COMPARISON_ID = "starved_vs_fed_0hpc"
HERE = Path(__file__).resolve().parent

SUPPLEMENT_URL = "https://www.ebi.ac.uk/europepmc/webservices/rest/PMC12803857/supplementaryFiles"
MEMBER = "mmc13.xlsx"
SHA256 = "451eb13e2da9db56ac984ee432632af2ced6ef9fc483183c62ba38c7c7f9c61d"

# Block header in Table S12 -> platform code used in the registry and sample sheet.
PLATFORMS = {
    "General Metabolites": "GM",
    "Nucleotides": "NUCs",
    "PUFAs": "PUFAs",
    "Krebs cycle": "TCA",
}
# The ten green rows, per the Methods (MES and the PUFA ISTD mix in the
# extraction solvent; d27-myristic acid for GC-MS).
INTERNAL_STANDARDS = {
    ("GM", "2-Morpholinoethansulfonic Acid"),
    ("NUCs", "MES"),
    ("PUFAs", "12-HETE-d8"), ("PUFAs", "6-keto-PGF1a-d4"), ("PUFAs", "LTB4 d4"),
    ("PUFAs", "OEA-d4"), ("PUFAs", "PAF-d4"), ("PUFAs", "PGD2-d4"), ("PUFAs", "TXB2-d4"),
    ("TCA", "d27-Myristic acid"),
}
# Named by the targeted dMRM / GC-MS method; the deposit confirms no identity
# against an authentic standard and gives no database identifier, so MSI
# level 2 (putatively annotated) rather than 1.
ANNOTATION_LEVEL = 2


def _sha256(path: Path) -> str:
    return hashlib.sha256(path.read_bytes()).hexdigest()


def download(staging: Path) -> Path:
    with urllib.request.urlopen(SUPPLEMENT_URL, timeout=120) as resp:
        archive = zipfile.ZipFile(io.BytesIO(resp.read()))
    out = staging / MEMBER
    out.write_bytes(archive.read(MEMBER))
    return out


def _is_green(cell) -> bool:
    fill = cell.fill
    return fill.fill_type == "solid" and fill.fgColor.type == "theme"


def parse_table_s12(path: Path) -> pd.DataFrame:
    ws = openpyxl.load_workbook(path).worksheets[0]
    rows, platform, samples = [], None, None
    for row in ws.iter_rows():
        name = row[0].value
        values = [c.value for c in row[1:]]
        if name in PLATFORMS:
            platform = PLATFORMS[name]
            samples = [v for v in values if v]
            continue
        if platform is None or name is None or values[0] == "Area":
            continue
        areas = values[: len(samples)]
        # A blank cell is missing (TXB3 has three, in 12 hpc samples), never 0.
        if not all(v is None or isinstance(v, (int, float)) for v in areas):
            raise SystemExit(f"{platform} {name!r}: non-numeric peak area")
        areas = [float("nan") if v is None else v for v in areas]
        rows.append({
            "feature_id": f"{platform}:{name}",
            "putative_metabolite_name": str(name).strip(),
            "annotation_confidence_level": ANNOTATION_LEVEL,
            "platform": platform,
            "is_internal_standard": _is_green(row[0]),
            **dict(zip(samples, areas)),
        })
    df = pd.DataFrame(rows)
    flagged = {(r.platform, r.putative_metabolite_name) for r in df.itertuples() if r.is_internal_standard}
    if flagged != INTERNAL_STANDARDS:
        raise SystemExit(f"green rows {sorted(flagged)} differ from the expected internal standards")
    if df["feature_id"].duplicated().any():
        raise SystemExit(f"duplicate features: {df.loc[df.feature_id.duplicated(), 'feature_id'].tolist()}")
    return df


def comparison_sheet() -> pd.DataFrame:
    """The 12 samples of starved_vs_fed_0hpc, with the tank as replicate unit.

    Replicate n is the tank (inferred, as for DUPERRET2025_STARVE_RNASEQ): each
    diet had three tanks holding both families.
    """
    sheet = pd.read_csv(HERE / "samplesheet.tsv", sep="\t", dtype=str)
    sheet = sheet[sheet["timepoint_hpc"] == "0"].copy()
    sheet["source_condition"] = sheet["condition"]
    sheet["condition"] = sheet["condition"].map({"fed": "control", "starved": "treatment"})
    sheet["tank"] = sheet["source_condition"] + "_t" + sheet["replicate"]
    return sheet[["sample_id", "condition", "source_condition", "family", "tank"]]


def main(argv=None) -> int:
    ap = argparse.ArgumentParser(description=__doc__.split("\n\n")[0])
    ap.add_argument("staging", type=Path, help="directory for the supplement and workflow inputs")
    ap.add_argument("--download", action="store_true", help=f"fetch {MEMBER} from Europe PMC first")
    args = ap.parse_args(argv)
    args.staging.mkdir(parents=True, exist_ok=True)

    source = download(args.staging) if args.download else args.staging / MEMBER
    if _sha256(source) != SHA256:
        raise SystemExit(f"{source}: SHA-256 {_sha256(source)} != expected {SHA256}")

    features = parse_table_s12(source)
    sheet = comparison_sheet()
    registered = pd.read_csv(HERE / "samplesheet.tsv", sep="\t", dtype=str)["sample_id"]
    sample_cols = [c for c in features.columns if c in set(registered)]
    if sorted(sample_cols) != sorted(registered):
        raise SystemExit("Table S12 sample columns do not match samplesheet.tsv")

    feat_path = args.staging / f"{STUDY_ID}_features.tsv"
    sheet_path = args.staging / f"{STUDY_ID}_{COMPARISON_ID}_samplesheet.tsv"
    features.to_csv(feat_path, sep="\t", index=False)
    sheet.to_csv(sheet_path, sep="\t", index=False)

    metabolites = features[~features["is_internal_standard"]]
    areas = metabolites[sample_cols]
    provenance = {
        "study_id": STUDY_ID,
        "date_generated": date.today().isoformat(),
        "converter": "data/studies/DUPERRET2025_POMS_STARVE/build_feature_table.py",
        "source": {
            "description": "Table S12. Raw metabolomic data: General Metabolites, Nucleotides, "
                           "PUFAs and Krebs cycle (peak areas)",
            "file": MEMBER,
            "url": SUPPLEMENT_URL,
            "article": "https://doi.org/10.1016/j.isci.2025.114333 (PMC12803857)",
            "license": "CC BY-NC-ND 4.0: not redistributed; fetched and checksum-verified by this script",
            "sha256": SHA256,
        },
        "outputs_sha256": {feat_path.name: _sha256(feat_path), sheet_path.name: _sha256(sheet_path)},
        "n_samples": len(sample_cols),
        "n_features": int(len(features)),
        "n_internal_standards": int(features["is_internal_standard"].sum()),
        "internal_standards": sorted(features.loc[features.is_internal_standard, "feature_id"]),
        "n_metabolites_by_platform": metabolites.groupby("platform").size().to_dict(),
        "n_metabolite_values_at_or_below_0.001": int((areas <= 0.001).sum().sum()),
        "n_metabolite_values": int(areas.size),
        "blank_cells": {fid: [c for c in sample_cols if pd.isna(row[c])]
                        for fid, row in metabolites.set_index("feature_id").iterrows()
                        if row[sample_cols].isna().any()},
        "annotation_confidence_level": ANNOTATION_LEVEL,
        "comparison": {
            "comparison_id": COMPARISON_ID,
            "samples": sheet["sample_id"].tolist(),
            "replicate_unit": "tank (= <diet>_t<replicate>; inferred, as in DUPERRET2025_STARVE_RNASEQ)",
        },
    }
    (HERE / "table_s12_provenance.json").write_text(json.dumps(provenance, indent=2) + "\n")
    print(f"{len(features)} features ({provenance['n_internal_standards']} internal standards) x "
          f"{len(sample_cols)} samples -> {feat_path}; {len(sheet)} samples -> {sheet_path}")
    return 0


if __name__ == "__main__":
    sys.exit(main())
